"""
docx_generator.py

Builds a clean, ATS-compliant .docx resume in a Harvard/professional style
from a structured JSON resume object. No tables, no text boxes, no columns —
only simple paragraph-based formatting that every ATS parser can read.
"""
import io
from typing import Optional
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _set_font(run, name="Calibri", size=10, bold=False, color: Optional[RGBColor] = None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def _add_section_heading(doc: Document, title: str):
    """Adds a section header with a bottom border rule (no text box needed)."""
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(5)
    para.paragraph_format.space_after = Pt(1.5)
    run = para.add_run(title.upper())
    _set_font(run, name="Calibri", size=10.5, bold=True, color=RGBColor(0x1A, 0x56, 0x8C))

    # Add bottom border to the paragraph
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A568C")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return para


def _add_bullet(doc: Document, text: str):
    para = doc.add_paragraph(style="List Bullet")
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(1)
    para.paragraph_format.line_spacing = 1.05
    run = para.add_run(text)
    _set_font(run, size=9.5)
    return para


def generate_docx_from_structured_resume(resume_data: dict) -> bytes:
    """
    Takes a structured resume dict and generates an ATS-compliant .docx
    closely matching Claude's 1-page Harvard standard.
    """
    doc = Document()

    # ── Page margins (0.5in top/bottom, 0.55in left/right for tight 1-page fit)
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    # ── Name (large, bold, centered)
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_para.paragraph_format.space_before = Pt(0)
    name_para.paragraph_format.space_after = Pt(1)
    name_run = name_para.add_run(resume_data.get("name", "").upper())
    _set_font(name_run, name="Calibri", size=17, bold=True)

    # ── Contact line (centered)
    contact_parts = []
    if resume_data.get("phone"): contact_parts.append(str(resume_data["phone"]).strip())
    if resume_data.get("email"): contact_parts.append(str(resume_data["email"]).strip())
    if resume_data.get("linkedin"): contact_parts.append(str(resume_data["linkedin"]).strip())
    if resume_data.get("github"): contact_parts.append(str(resume_data["github"]).strip())

    contact_para = doc.add_paragraph()
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_para.paragraph_format.space_before = Pt(0)
    contact_para.paragraph_format.space_after = Pt(3)
    contact_run = contact_para.add_run(" | ".join(contact_parts))
    _set_font(contact_run, size=9.5, color=RGBColor(0x33, 0x33, 0x33))

    # ── Summary (optional, dense 4-sentence authoritative profile)
    if resume_data.get("summary"):
        _add_section_heading(doc, "Summary")
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(resume_data["summary"].strip())
        _set_font(r, size=9.5)

    # ── Experience
    experience = resume_data.get("experience", [])
    if experience:
        _add_section_heading(doc, "Experience")
        for job in experience:
            role_para = doc.add_paragraph()
            role_para.paragraph_format.space_before = Pt(3)
            role_para.paragraph_format.space_after = Pt(0)
            
            title = (job.get('title') or '').strip()
            company = (job.get('company') or '').strip()
            header_str = f"{title} — {company}" if title and company else (title or company)
            title_run = role_para.add_run(header_str)
            _set_font(title_run, size=10, bold=True)

            # Location + dates (smaller, grey, defensive against "None")
            loc = (job.get('location') or '').strip()
            if loc.lower() in ("none", "null", ""):
                loc = ""
            dates = (job.get('dates') or '').strip()
            if dates.lower() in ("none", "null"):
                dates = ""

            meta_text = f"{loc} | {dates}" if loc and dates else (dates or loc)
            if meta_text:
                meta_para = doc.add_paragraph()
                meta_para.paragraph_format.space_before = Pt(0)
                meta_para.paragraph_format.space_after = Pt(1.5)
                meta_run = meta_para.add_run(meta_text)
                _set_font(meta_run, size=9, color=RGBColor(0x55, 0x55, 0x55))

            for bullet in job.get("bullets", []):
                if bullet and bullet.strip():
                    _add_bullet(doc, bullet.strip())

    # ── Projects
    projects = resume_data.get("projects", [])
    if projects:
        _add_section_heading(doc, "Projects")
        for proj in projects:
            proj_para = doc.add_paragraph()
            proj_para.paragraph_format.space_before = Pt(3)
            proj_para.paragraph_format.space_after = Pt(0.5)
            
            proj_name = (proj.get("name") or "").strip()
            proj_name_run = proj_para.add_run(proj_name)
            _set_font(proj_name_run, size=10, bold=True)
            
            # Repo link or tech label beside name
            tech = (proj.get("tech") or proj.get("link") or "").strip()
            if tech and tech.lower() not in ("none", "null"):
                tech_run = proj_para.add_run(f" | {tech}")
                _set_font(tech_run, size=9, color=RGBColor(0x55, 0x55, 0x55))
                
            for bullet in proj.get("bullets", []):
                if bullet and bullet.strip():
                    _add_bullet(doc, bullet.strip())

    # ── Education
    education = resume_data.get("education", [])
    if education:
        _add_section_heading(doc, "Education")
        for edu in education:
            edu_para = doc.add_paragraph()
            edu_para.paragraph_format.space_before = Pt(3)
            edu_para.paragraph_format.space_after = Pt(0.5)
            
            deg = (edu.get('degree') or '').strip()
            inst = (edu.get('institution') or '').strip()
            edu_str = f"{deg} — {inst}" if deg and inst else (deg or inst)
            deg_run = edu_para.add_run(edu_str)
            _set_font(deg_run, size=10, bold=True)
            
            dates = (edu.get("dates") or "").strip()
            gpa = (edu.get("gpa") or "").strip()
            meta_parts = []
            if dates and dates.lower() not in ("none", "null"):
                meta_parts.append(dates)
            if gpa and gpa.lower() not in ("none", "null"):
                meta_parts.append(f"GPA: {gpa}")
                
            if meta_parts:
                edu_meta = doc.add_paragraph()
                edu_meta.paragraph_format.space_before = Pt(0)
                edu_meta.paragraph_format.space_after = Pt(1.5)
                meta_run = edu_meta.add_run(" | ".join(meta_parts))
                _set_font(meta_run, size=9, color=RGBColor(0x55, 0x55, 0x55))

    # ── Technical Skills
    skills = resume_data.get("skills", {})
    if skills:
        _add_section_heading(doc, "Technical Skills")
        skill_lines = []
        if skills.get("languages"):   skill_lines.append(("Languages", skills["languages"]))
        if skills.get("frameworks"):  skill_lines.append(("Frameworks", skills["frameworks"]))
        if skills.get("tools"):       skill_lines.append(("Tools", skills["tools"]))
        if skills.get("databases"):   skill_lines.append(("Databases", skills["databases"]))
        for label, value in skill_lines:
            if value and str(value).strip() and str(value).strip().lower() not in ("none", "null"):
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(0.5)
                p.paragraph_format.space_after = Pt(0.5)
                label_run = p.add_run(f"{label}: ")
                _set_font(label_run, size=9.5, bold=True)
                val_run = p.add_run(str(value).strip())
                _set_font(val_run, size=9.5)

    # ── Save to bytes buffer
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()
